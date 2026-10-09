"""Resume parser tests (task 9.1) — format detection, honest extraction, loud failures."""

import json

import pytest

from engine.ats.parser import (
    PARSER_VERSION,
    detect_format,
    extract_docx,
    extract_pdf,
    parse_resume,
)
from engine.errors import VantiaError

MARKDOWN = """# Ada Okafor
Backend Engineer
Location: Lagos, Nigeria
ada@example.com | +234 801 234 5678

## Summary
Ships reliable services.

## Skills
- Python, Kubernetes
- PostgreSQL

## Experience
SRE, Acme (Jan 2022 - Present)
- Shipped the migration, 20% faster

## Education
BSc Computer Science, UNN (2020)
"""

JSON_PROFILE = {
    "name": "Ada Okafor",
    "title": "Backend Engineer",
    "email": "ada@example.com",
    "skills": ["Python", "Kubernetes"],
    "experience": [
        {
            "role": "SRE",
            "org": "Acme",
            "start": "2022",
            "end": "2024",
            "summary": ["Automated deploys"],
        }
    ],
    "education": ["BSc Computer Science"],
}


def test_detect_format_magic_and_hints():
    assert detect_format(b"%PDF-1.4...") == "pdf"
    assert detect_format(b"PK\x03\x04rest") == "docx"
    assert detect_format(b"anything", hint="resume.md") == "md"
    assert detect_format(b"anything", hint="resume.markdown") == "md"
    assert detect_format(b"anything", hint="profile.json") == "json"
    assert detect_format(b"anything", hint="resume.txt") == "txt"
    assert detect_format(b"plain words") == "txt"


def test_json_profile_is_high_confidence_and_survives_missing_keys():
    # regression: profiles without languages/certifications used to crash `_strings`
    resume = parse_resume(JSON_PROFILE)
    assert resume.source_format == "json"
    assert resume.extraction_confidence == "high"
    assert resume.name == "Ada Okafor" and resume.title == "Backend Engineer"
    assert resume.skills == ["Python", "Kubernetes"]
    assert resume.languages == [] and resume.certifications == []
    assert resume.metadata["parser_version"] == PARSER_VERSION
    assert resume.metadata["source_format"] == "json"
    assert resume.years_experience == pytest.approx(2.0, abs=0.1)  # 2022 → 2024


def test_json_profile_accepts_comma_strings_skill_objects_and_missing_name():
    resume = parse_resume(
        {
            "skills": "Python, Go",
            "languages": [{"name": "English"}, "Igbo"],
            "education": [{"degree": "BSc CS", "school": "UNN", "year": "2020"}],
        }
    )
    assert resume.skills == ["Python", "Go"]
    assert resume.languages == ["English", "Igbo"]
    assert resume.education[0]["degree"] == "BSc CS"
    assert resume.extraction_confidence == "high"
    assert any("name" in w for w in resume.parse_warnings)


def test_json_bytes_and_invalid_json():
    resume = parse_resume(json.dumps(JSON_PROFILE).encode(), content_type="application/json")
    assert resume.source_format == "json" and resume.name == "Ada Okafor"
    with pytest.raises(VantiaError) as exc:
        parse_resume(b"{not json", content_type="application/json")
    assert exc.value.code == "profile_parse_error"


def test_missing_file_and_unsupported_binary_refused():
    with pytest.raises(VantiaError) as exc:
        parse_resume("/nonexistent/resume.pdf")
    assert exc.value.code == "profile_file_missing"
    with pytest.raises(VantiaError) as exc:
        parse_resume(b"\xff\xfe\x00\x00garbage", content_type="text/plain")
    assert exc.value.code == "profile_format_unsupported"


def test_empty_resume_refused_not_scored_as_zero():
    with pytest.raises(VantiaError) as exc:
        parse_resume(b"   \n  ", content_type="text/plain")
    assert exc.value.code == "resume_parse_error"


def test_markdown_sections_contacts_and_confidence():
    resume = parse_resume(MARKDOWN.encode(), content_type="text/markdown")
    assert resume.source_format == "md"
    assert resume.name == "Ada Okafor" and resume.title == "Backend Engineer"
    assert resume.email == "ada@example.com"
    assert "Lagos" in resume.location
    assert {"summary", "skills", "experience", "education"} <= set(resume.sections)
    assert "Python" in resume.skills and "PostgreSQL" in resume.skills
    entry = resume.experience[0]
    assert entry["role"] == "SRE" and entry["org"] == "Acme"
    assert entry["start"] == "2022-01" and entry["end"] == "present"
    assert any("20% faster" in line for line in entry["summary"])
    assert any("BSc" in e["degree"] for e in resume.education)
    # email found → not downgraded to low
    assert resume.extraction_confidence == "partial"
    assert not any("email" in w for w in resume.parse_warnings)


def test_missing_email_lowers_confidence_and_warns():
    resume = parse_resume(b"Jane Doe\n\n## Skills\n- Go\n", content_type="text/markdown")
    assert resume.extraction_confidence == "low"
    assert any("email" in w for w in resume.parse_warnings)


def test_pdf_extraction_injected_and_loud_paths():
    text = extract_pdf(b"%PDF-fake", extractor=lambda _d: "Ada\nSkills: Python")
    assert "Ada" in text
    with pytest.raises(VantiaError) as exc:  # empty extraction is never an empty resume
        extract_pdf(b"%PDF", extractor=lambda _d: "  ")
    assert exc.value.code == "resume_parse_error"

    def boom(_data: bytes) -> str:
        raise ModuleNotFoundError("pypdf extractor missing", name="pypdf")

    with pytest.raises(VantiaError) as exc:  # missing extra → install hint, not a crash
        extract_pdf(b"%PDF", extractor=boom)
    assert exc.value.code == "resume_extractor_unavailable"
    assert "ats" in str(exc.value) or "extra" in str(exc.value)

    with pytest.raises(VantiaError) as exc:  # corrupt PDF through the real pypdf path
        extract_pdf(b"%PDF-1.4 not really a pdf")
    assert exc.value.code == "resume_parse_error"


def test_docx_roundtrip_and_loud_failure(tmp_path):
    import docx

    doc = docx.Document()
    doc.add_paragraph("Ada Okafor")
    doc.add_paragraph("Backend Engineer")
    doc.add_paragraph("Skills: Python, Kubernetes")
    path = tmp_path / "resume.docx"
    doc.save(str(path))
    resume = parse_resume(str(path))
    assert resume.source_format == "docx"
    assert "Ada Okafor" in resume.raw_text and "Kubernetes" in resume.raw_text

    with pytest.raises(VantiaError) as exc:  # corrupt DOCX refuses loudly
        extract_docx(b"PK\x03\x04 not a real docx")
    assert exc.value.code == "resume_parse_error"


def test_years_experience_dated_versus_undated():
    dated = parse_resume(
        {"experience": [{"role": "A"}, {"role": "B", "start": "2020", "end": "2023"}]}
    )
    assert dated.years_experience == pytest.approx(3.0, abs=0.01)
    assert dated.has_dates is True
    undated = parse_resume({"experience": [{"role": "A"}]})
    assert undated.years_experience == 0.0
    assert undated.has_dates is False


def test_to_dict_roundtrip_shape():
    resume = parse_resume(JSON_PROFILE)
    data = resume.to_dict()
    assert {"name", "skills", "sections", "metadata", "parse_warnings"} <= set(data)
    assert data["metadata"]["source_format"] == "json"
    assert data["sections"]["skills"]
